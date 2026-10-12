import argparse
import json
from itertools import zip_longest
from pathlib import Path

from openpyxl import load_workbook


def _cell_value(value: object) -> str:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return repr(value)


def _walk_json(left: object, right: object, path: str, differences: list[str]) -> None:
    if isinstance(left, dict) and isinstance(right, dict):
        for key in sorted(left.keys() | right.keys()):
            _walk_json(left.get(key), right.get(key), f"{path}.{key}", differences)
        return
    if isinstance(left, list) and isinstance(right, list):
        for index, (lhs, rhs) in enumerate(zip_longest(left, right)):
            _walk_json(lhs, rhs, f"{path}[{index}]", differences)
        return
    if left != right:
        differences.append(f"{path}: {left!r} != {right!r}")


def _compare_json(expected: Path, actual: Path) -> list[str]:
    before = json.loads(expected.read_text(encoding="utf-8"))
    after = json.loads(actual.read_text(encoding="utf-8"))
    if before == after:
        return []
    differences = []

    _walk_json(before, after, expected.name, differences)
    return differences


def _compare_excel(expected: Path, actual: Path) -> list[str]:
    before = load_workbook(expected, read_only=True, data_only=False)
    after = load_workbook(actual, read_only=True, data_only=False)
    differences = []
    try:
        if before.sheetnames != after.sheetnames:
            differences.append(f"{expected.name}: sheets differ")
            return differences
        for sheet_name in before.sheetnames:
            differences.extend(_compare_sheet(
                expected.name, sheet_name, before[sheet_name], after[sheet_name]
            ))
    finally:
        before.close()
        after.close()
    return differences


def _compare_sheet(filename: str, sheet_name: str, left_sheet, right_sheet) -> list[str]:
    differences = []
    rows = zip_longest(
        left_sheet.iter_rows(values_only=True), right_sheet.iter_rows(values_only=True)
    )
    for row_index, (left, right) in enumerate(rows, 1):
        differences.extend(_compare_excel_row(filename, sheet_name, row_index, left, right))
    return differences


def _compare_excel_row(filename: str, sheet_name: str, row_index: int,
                       left: tuple | None, right: tuple | None) -> list[str]:
    left, right = left or (), right or ()
    if _contains_generated_timestamp(left, right):
        return []
    differences = []
    for column_index, (left_value, right_value) in enumerate(zip_longest(left, right), 1):
        if left_value != right_value:
            differences.append(
                f"{filename}/{sheet_name}/R{row_index}C{column_index}: "
                f"{_cell_value(left_value)} != {_cell_value(right_value)}"
            )
    return differences


def _contains_generated_timestamp(left: tuple, right: tuple) -> bool:
    return any(
        isinstance(value, str) and "Documento gerado em:" in value
        for value in left + right
    )


def compare(expected: Path, actual: Path) -> list[str]:
    differences = []
    expected_periods = {path.stem for path in expected.glob("????-??.json")}
    actual_periods = {path.stem for path in actual.glob("????-??.json")}
    if expected_periods != actual_periods:
        differences.append(f"Different periods: {expected_periods ^ actual_periods}")
    for period in sorted(expected_periods & actual_periods):
        differences.extend(_compare_json(expected / f"{period}.json", actual / f"{period}.json"))
        differences.extend(_compare_excel(expected / f"{period}.xlsx", actual / f"{period}.xlsx"))
    return differences


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare two apuration golden directories")
    parser.add_argument("expected", type=Path)
    parser.add_argument("actual", type=Path)
    args = parser.parse_args()
    differences = compare(args.expected, args.actual)
    for difference in differences[:100]:
        print(difference)
    print(f"Differences: {len(differences)}")
    if differences:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
