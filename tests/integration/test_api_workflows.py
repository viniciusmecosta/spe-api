import asyncio
import json
import os
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import httpx
import pytest
import pytest_asyncio
from api_workflow_support import (
    DAY,
    PASSWORD,
    PERIOD,
    ApiWorkflow,
    fixed_clock,
    logo_bytes,
    seed_root,
)
from openpyxl import load_workbook
from reportlab.pdfgen import canvas
from sqlalchemy import func, select

from app.features.adjustments.adjustment_models import AdjustmentRequest
from app.features.daily_summaries.daily_summary_models import DailySummary
from app.features.system import job_queue
from app.shared.enums import AdjustmentType

pytestmark = pytest.mark.asyncio


@pytest_asyncio.fixture
async def api(pg_session_factory, monkeypatch):
    from app.main import app
    from app.shared.trusted_time_service import trusted_time_service

    await seed_root(pg_session_factory)
    monkeypatch.setattr(trusted_time_service, "get_trusted_time", fixed_clock)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app, raise_app_exceptions=False), base_url="http://test"
    ) as client:
        flow = ApiWorkflow(client)
        flow.headers = await flow.login("qa_root")
        yield flow
    evidence = os.getenv("SPE_ENDPOINT_EVIDENCE")
    if evidence:
        await asyncio.to_thread(write_evidence, evidence, flow.calls)


def write_evidence(path, calls):
    with Path(path).open("a") as output:
        output.writelines(json.dumps({"method": method, "path": path, "status": status}) + "\n"
                          for method, path, status in calls)


async def test_user_company_and_hardware_administration(api):
    await api.request("GET", "/health", headers={})
    await api.request("GET", "/ready", headers={})
    await api.request("GET", "/operations")
    await api.request("GET", "/auth/me")
    await api.request("POST", "/time-records/entry", 201, headers={**api.headers, "X-Platform": "desktop"})
    await api.request("POST", "/time-records/exit", 201, headers={**api.headers, "X-Platform": "desktop"})
    await api.request("GET", "/users/me")
    await api.request("PUT", "/users/me", json={"name": "Gestor de Teste"})
    user_id = await api.create_employee()
    await api.request("GET", "/users/", params={"limit": 1, "order_by": "id"})
    await api.request("GET", f"/users/{user_id}")
    await api.request("PUT", f"/users/{user_id}", json={"name": "Colaborador Beta"})
    await api.request("POST", "/users/", 403, headers=api.employee_headers,
                      json={"username": "forbidden", "password": PASSWORD})
    company = await api.json("POST", "/companies/", json={
        "name": "Empresa Sintetica", "cnpj": "12345678000195", "address": "Endereco de teste",
    })
    await api.request("GET", "/companies/")
    await api.request("PUT", "/companies/", json={"phone": "0000000000"})
    await api.request("POST", "/companies/logo", files={"file": ("logo.png", logo_bytes(), "image/png")})
    printer = await api.json("POST", "/printers/", 201, json={
        "name": "Impressora Sintetica", "address": "127.0.0.1", "company_id": company["id"], "status": False,
    })
    await api.request("GET", "/printers/")
    await api.request("GET", f"/printers/{printer['id']}")
    await api.request("PATCH", f"/printers/{printer['id']}", json={"paper_width": 58})
    await api.request("DELETE", f"/printers/{printer['id']}", 204)
    credential = await api.json("POST", "/device-credentials/", 201, json={
        "name": "Dispositivo Sintetico", "key_type": "DEVICE", "api_key": "synthetic-device-key",
    })
    device_headers = {"X-API-KEY": "synthetic-device-key"}
    await api.request("GET", "/device-credentials/")
    await api.request("PUT", f"/device-credentials/{credential['id']}", json={"name": "Dispositivo Atualizado"})
    await api.request("GET", "/biometrics/available-sensor-indices")
    await api.request("GET", "/devices/time", headers=device_headers)
    bootstrap = await api.json("POST", "/devices/verify-manager", headers=device_headers, json={"sensor_index": 11})
    assert bootstrap["is_allowed"] is True
    await api.request("PUT", "/users/1", json={"biometrics": [{"sensor_index": 12, "template_data": "SYNTHETIC_MANAGER", "finger_id": 0}]})
    denied = await api.json("POST", "/devices/verify-manager", headers=device_headers, json={"sensor_index": 11})
    assert denied["is_allowed"] is False
    allowed = await api.json("POST", "/devices/verify-manager", headers=device_headers, json={"sensor_index": 12})
    assert allowed["is_allowed"] is True
    await api.request("POST", "/devices/punch", headers=device_headers, json={"sensor_index": 11})
    await api.request("POST", "/firmware/upload", 201, files={"file": ("firmware.bin", b"synthetic-firmware", "application/octet-stream")}, data={"version": "v1.0.0"})
    await api.request("GET", "/firmware/")
    await api.request("PUT", "/firmware/v1.0.0", files={"file": ("firmware.bin", b"updated-firmware", "application/octet-stream")})
    await api.request("GET", "/firmware/check", headers=device_headers)
    download = await api.request("GET", "/firmware/download", headers=device_headers, params={"version": "v1.0.0"})
    assert download.content == b"updated-firmware"
    await api.request("DELETE", f"/device-credentials/{credential['id']}")
    await api.request("GET", "/devices/time", 401, headers=device_headers)


async def test_schedule_records_reports_and_payroll(api):
    user_id = await api.create_employee()
    await api.json("POST", "/companies/", json={"name": "Empresa Sintetica", "cnpj": "12345678000195", "address": "Teste"})
    schedule = {"valid_from": "2026-09-01", "valid_until": "2026-09-30", "users": [{"user_id": user_id, "schedules": [
        {"day_of_week": 1, "daily_hours": 8, "entry_1": "08:00", "exit_1": "12:00", "entry_2": "13:00", "exit_2": "17:00"},
    ]}]}
    await api.request("POST", "/schedules/bulk", json=schedule)
    await api.request("GET", "/schedules/bulk", params=PERIOD)
    period_path = "/schedules/bulk/2026-09-01/2026-09-30"
    await api.request("GET", period_path)
    await api.request("PUT", period_path, json=schedule)
    holiday = await api.json("POST", "/holidays/", 201, json={"date": "2026-09-28", "name": "Feriado Sintetico"})
    await api.request("GET", "/holidays/", params=PERIOD)
    await api.request("DELETE", f"/holidays/{holiday['id']}")
    entry = await api.add_punch(user_id, "ENTRY", "08:00")
    exit_record = await api.add_punch(user_id, "EXIT", "16:00")
    edited = await api.json("PUT", f"/time-records/admin/{exit_record['id']}", json={
        "record_datetime": f"{DAY}T17:00:00-03:00", "edit_justification": "Correcao sintetica",
    })
    await api.request("GET", "/time-records/admin/list", params={"user_id": user_id, "start_date": DAY, "end_date": DAY})
    await api.request("GET", f"/time-records/{edited['id']}/timeline")
    await api.request("GET", "/time-records/my", headers=api.employee_headers)
    await api.request("GET", f"/time-records/receipt/{entry['short_id']}", headers=api.employee_headers)
    receipt = await api.request("GET", f"/time-records/receipt/{entry['short_id']}/pdf", headers=api.employee_headers)
    assert receipt.content.startswith(b"%PDF")
    for path in ["/reports/dashboard", "/dashboard/manager", "/reports/team-hours", f"/reports/user/{user_id}",
                 f"/reports/history/user/{user_id}", "/anomalies/all", f"/anomalies/user/{user_id}"]:
        await api.request("GET", path, params=PERIOD)
    for path in ["/reports/my/dashboard", "/reports/history/me"]:
        await api.request("GET", path, headers=api.employee_headers, params=PERIOD)
    excel = await api.request("GET", "/reports/export/excel", params=PERIOD)
    workbook = load_workbook(BytesIO(excel.content), data_only=True)
    assert len(workbook.sheetnames) >= 2
    pdf = await api.request("GET", f"/timesheets/user/{user_id}/pdf", params=PERIOD)
    assert pdf.content.startswith(b"%PDF")
    archive = await api.request("GET", "/timesheets/all/pdf", params=PERIOD)
    assert archive.content.startswith(b"PK")
    closure = await api.json("POST", "/payroll/close", json=PERIOD)
    await api.request("GET", "/payroll/", params=PERIOD)
    await api.request("POST", "/adjustments/admin/waive", 400, json={
        "user_id": user_id, "target_date": DAY, "amount_hours": 1, "reason_text": "Mes fechado",
    })
    await api.request("GET", f"/payroll/{closure['id']}/download-report")
    await api.request("POST", f"/payroll/{closure['id']}/legacy-report",
                      files={"file": ("legacy.xlsx", excel.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    await api.request("POST", "/payroll/reopen", json={**PERIOD, "observation": "Reabertura sintetica"})
    await api.request("PUT", f"/time-records/{edited['id']}/toggle", headers=api.employee_headers)
    await api.request("DELETE", f"/time-records/admin/{entry['id']}", json={"edit_justification": "Exclusao sintetica"})
    await api.request("DELETE", period_path)


async def test_adjustment_and_durable_job_workflows(api, pg_session_factory):
    user_id = await api.create_employee()
    request = await api.json("POST", "/adjustments/", 201, headers=api.employee_headers, json={
        "adjustment_type": "WAIVER", "target_date": DAY, "amount_hours": 1, "reason_text": "Atestado sintetico",
    })
    pdf = BytesIO()
    document = canvas.Canvas(pdf)
    document.drawString(10, 10, "Documento sintetico")
    document.save()
    await api.request("POST", f"/adjustments/{request['id']}/attachments", 201, headers=api.employee_headers,
                      files={"file": ("atestado.pdf", pdf.getvalue(), "application/pdf")})
    download = await api.request("GET", f"/adjustments/{request['id']}/download", headers=api.employee_headers)
    assert download.content == pdf.getvalue()
    await api.request("GET", "/adjustments/my", headers=api.employee_headers)
    await api.request("GET", "/adjustments/", params=PERIOD)
    await api.request("PUT", f"/adjustments/{request['id']}/approve", json={"comment": "Validado"})
    await api.request("PUT", f"/adjustments/{request['id']}/reject", json={"comment": "Rejeicao sintetica"})
    await api.request("PUT", f"/adjustments/admin/{request['id']}/revert-status", json={"status": "PENDING", "comment": "Revisao sintetica"})
    await api.request("DELETE", f"/adjustments/{request['id']}", params={"reason": "Exclusao sintetica"})
    waiver = await api.json("POST", "/adjustments/admin/waive", 201, json={
        "user_id": user_id, "target_date": DAY, "amount_hours": 2, "reason_text": "Abono administrativo sintetico",
    })
    assert waiver["status"] == "APPROVED"
    async with pg_session_factory() as session:
        count = await session.scalar(select(func.count()).select_from(AdjustmentRequest).where(
            AdjustmentRequest.user_id == user_id, AdjustmentRequest.adjustment_type == AdjustmentType.WAIVER,
            AdjustmentRequest.deleted_at.is_(None),
        ))
        assert count == 1
    await api.request("DELETE", f"/adjustments/admin/{waiver['id']}", params={"reason": "Exclusao sintetica"})
    jobs = []
    for path in ["/backup/trigger", "/telegram/manual-backup"]:
        jobs.append(await api.json("POST", path))
    jobs.append(await api.json("POST", "/telegram/manual-report", params={"start_date": DAY, "end_date": DAY}))
    jobs.append(await api.json("POST", "/daily-summaries/admin/reprocess", json={"start_date": DAY, "end_date": DAY, "user_ids": [user_id]}))
    for job in jobs:
        result = await api.json("GET", f"/routine-logs/jobs/{job['job_id']}")
        assert result["status"] == "PENDING"
    await api.request("POST", f"/routine-logs/jobs/{jobs[0]['job_id']}/retry", 409)
    claimed = await job_queue.claim()
    assert claimed["id"] == int(jobs[0]["job_id"])
    assert await job_queue.checkpoint(claimed, "FAILED")
    await api.request("POST", f"/routine-logs/jobs/{jobs[0]['job_id']}/retry")
    retried = await api.json("GET", f"/routine-logs/jobs/{jobs[0]['job_id']}")
    assert retried["status"] == "PENDING"
    await api.request("GET", "/routine-logs/")
    await api.request("GET", "/audit/")
    await api.request("GET", "/operations")


async def test_all_private_routes_reject_unauthenticated_access(api):
    import re

    from app.main import app

    public = {("GET", "/api/v1/health"): 200, ("GET", "/api/v1/ready"): 200,
              ("POST", "/api/v1/auth/login"): 422}
    for path, methods in app.openapi()["paths"].items():
        concrete = path.replace("{valid_from}", "2026-09-01").replace("{valid_until}", "2026-09-30")
        concrete = re.sub(r"\{[^}]+\}", "999999", concrete)
        for method in methods:
            if method.upper() not in {"GET", "POST", "PUT", "PATCH", "DELETE"}:
                continue
            expected = public.get((method.upper(), path), 401)
            await api.request(method.upper(), concrete.removeprefix("/api/v1"), expected, headers={})


async def test_apuration_history_and_excel_preserve_gross_and_partial_excess(api, pg_session_factory):
    user_id = await api.create_employee()
    await api.request("POST", "/companies/", json={"name": "Empresa Sintetica", "cnpj": "12345678000195", "address": "Teste"})
    await api.request("POST", "/schedules/bulk", json={
        "valid_from": "2026-09-01", "valid_until": "2026-09-30", "users": [{"user_id": user_id, "schedules": [{
            "day_of_week": 1, "daily_hours": 8, "entry_1": "08:00", "exit_1": "12:00",
            "entry_2": "13:00", "exit_2": "17:00", "is_daily_excess_enabled": True,
        }]}],
    })
    await api.add_punch(user_id, "ENTRY", "08:00")
    await api.add_punch(user_id, "EXIT", "17:00")
    adjustments = await api.json("GET", "/adjustments/", params=PERIOD)
    excess = next(item for item in adjustments if item["adjustment_type"] == "DAILY_EXCESS")
    await api.request("PUT", f"/adjustments/{excess['id']}/approve", json={"approved_amount_hours": 0.5})
    await api.request("POST", "/adjustments/admin/waive", 201, json={
        "user_id": user_id, "target_date": "2026-09-22", "amount_hours": 2, "reason_text": "Abono sintetico",
    })
    await api.request("POST", "/holidays/", 201, json={"date": "2026-09-15", "name": "Feriado sintetico"})
    history = await api.json("GET", f"/reports/history/user/{user_id}", params=PERIOD)
    assert history["total_worked_time"] == "11:00"
    assert history["total_accounted_time"] == "10:30"
    day = next(item for item in history["days"] if item["date"] == DAY)
    assert day["daily_excess"]["approved_minutes"] == 30
    assert day["daily_excess"]["unapproved_minutes"] == 30
    async with pg_session_factory() as session:
        rows = list((await session.scalars(select(DailySummary).where(DailySummary.user_id == user_id))).all())
        saved = next(row for row in rows if row.apuration_date == date.fromisoformat(DAY))
        assert (saved.worked_minutes, saved.accounted_minutes, saved.authorized_excess_minutes) == (540, 510, 30)
        waived = next(row for row in rows if row.apuration_date == date(2026, 9, 22))
        assert (waived.worked_minutes, waived.waiver_minutes, waived.missing_minutes) == (120, 120, 360)
        holiday = next(row for row in rows if row.apuration_date == date(2026, 9, 15))
        assert holiday.expected_minutes == 0
    excel = await api.request("GET", "/reports/export/excel", params=PERIOD)
    workbook = load_workbook(BytesIO(excel.content), data_only=True)
    sheet = workbook["Colaborador Alfa"]
    row = next(row for row in sheet.iter_rows() if row[0].value == "29/09/26")
    assert row[18].value == timedelta(hours=9)
    assert row[22].value == timedelta(hours=8, minutes=30)
